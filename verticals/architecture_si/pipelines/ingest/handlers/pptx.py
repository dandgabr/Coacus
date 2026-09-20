"""PPTX to Markdown slides (text frames and tables)."""

from __future__ import annotations

from pathlib import Path

from engine.dispatcher import register_converter

from verticals.architecture_si.pipelines.ingest.handlers import _common as common


def convert_pptx_to_md(pptx_path: str) -> str:
    if Path(pptx_path).suffix.lower() == ".ppt":
        raise RuntimeError(
            "legacy .ppt (OLE) is not supported by python-pptx; convert it to "
            ".pptx first (e.g. LibreOffice)"
        )
    try:
        import pptx
    except ImportError as exc:  # pragma: no cover - optional dependency
        raise ImportError("python-pptx is required: pip install python-pptx") from exc

    prs = pptx.Presentation(pptx_path)
    md_lines = [f"# Presentation: {Path(pptx_path).stem}\n"]
    for idx, slide in enumerate(prs.slides, start=1):
        md_lines.append(f"\n---\n\n## Slide {idx}")
        for shape in slide.shapes:
            if shape.has_text_frame:
                for paragraph in shape.text_frame.paragraphs:
                    text = paragraph.text.strip()
                    if not text:
                        continue
                    if shape == slide.shapes[0] and idx > 1:
                        md_lines.append(f"### {text}\n")
                    else:
                        md_lines.append(f"- {text}")
            elif shape.has_table:
                rows = [
                    [cell.text.strip().replace("\n", " ") for cell in row.cells]
                    for row in shape.table.rows
                ]
                if rows:
                    md_lines.append("\n" + common.format_markdown_table(rows) + "\n")
    return "\n".join(md_lines)


@register_converter(".pptx", ".ppt")
def handle_pptx(input_path: str, output_path: str | None = None) -> str:
    return common.write_markdown(convert_pptx_to_md(input_path), output_path, input_path)
