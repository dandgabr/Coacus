"""DOCX to Markdown (headings, paragraphs, lists, tables)."""

from __future__ import annotations

from pathlib import Path

from engine.dispatcher import register_converter

from verticals.architecture_si.pipelines.ingest.handlers import _common as common


def convert_docx_to_md(docx_path: str) -> str:
    if Path(docx_path).suffix.lower() == ".doc":
        raise RuntimeError(
            "legacy .doc (OLE) is not supported by python-docx; convert it to "
            ".docx first (e.g. LibreOffice)"
        )
    try:
        import docx
    except ImportError as exc:  # pragma: no cover - optional dependency
        raise ImportError("python-docx is required: pip install python-docx") from exc

    doc = docx.Document(docx_path)
    md_lines = [f"# {Path(docx_path).stem}\n"]
    for para in doc.paragraphs:
        text = para.text.strip()
        if not text:
            continue
        style = para.style.name.lower()
        if "heading 1" in style:
            md_lines.append(f"\n# {text}\n")
        elif "heading 2" in style:
            md_lines.append(f"\n## {text}\n")
        elif "heading 3" in style:
            md_lines.append(f"\n### {text}\n")
        elif "list" in style or text.startswith(("-", "*", "•")):
            md_lines.append(f"- {text.lstrip('-*• ')}")
        else:
            md_lines.append(f"{text}\n")
    for table in doc.tables:
        rows = [
            [cell.text.strip().replace("\n", " ") for cell in row.cells]
            for row in table.rows
        ]
        if rows:
            md_lines.append("\n" + common.format_markdown_table(rows) + "\n")
    return "\n".join(md_lines)


@register_converter(".docx", ".doc")
def handle_docx(input_path: str, output_path: str | None = None) -> str:
    return common.write_markdown(convert_docx_to_md(input_path), output_path, input_path)
